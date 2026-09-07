from docx import Document
from docx.shared import Pt, RGBColor, Inches, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

doc = Document()

# ── Page margins ──────────────────────────────────────────────────────────────
for section in doc.sections:
    section.top_margin    = Cm(2.0)
    section.bottom_margin = Cm(2.0)
    section.left_margin   = Cm(2.5)
    section.right_margin  = Cm(2.5)

# ── Colour palette ────────────────────────────────────────────────────────────
SAFFRON = RGBColor(0xFF, 0x99, 0x33)
NAVY    = RGBColor(0x0D, 0x0D, 0x1A)
DGREY   = RGBColor(0x44, 0x44, 0x55)
BLACK   = RGBColor(0x1A, 0x1A, 0x1A)

# ── Helpers ───────────────────────────────────────────────────────────────────

def set_font(run, size, bold=False, color=BLACK, name="Calibri"):
    run.font.name = name
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color

def shade_paragraph(para, hex_fill):
    """Apply a background shade to a paragraph."""
    pPr = para._p.get_or_add_pPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), hex_fill)
    pPr.append(shd)

def add_section_heading(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after  = Pt(2)
    shade_paragraph(p, '0D0D1A')
    run = p.add_run(f"  {text}")
    set_font(run, 13, bold=True, color=RGBColor(0xFF,0xFF,0xFF), name="Calibri")
    return p

def add_field(doc, label, value, placeholder=False):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after  = Pt(3)
    p.paragraph_format.left_indent  = Inches(0.2)
    r_label = p.add_run(f"{label}:  ")
    set_font(r_label, 11, bold=True, color=DGREY)
    r_value = p.add_run(value)
    if placeholder:
        set_font(r_value, 11, bold=False, color=RGBColor(0xAA, 0x66, 0x00))
    else:
        set_font(r_value, 11, bold=False, color=BLACK)
    return p

def add_bullet(doc, text, placeholder=False):
    p = doc.add_paragraph(style='List Bullet')
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after  = Pt(2)
    p.paragraph_format.left_indent  = Inches(0.4)
    run = p.add_run(text)
    if placeholder:
        set_font(run, 11, color=RGBColor(0xAA, 0x66, 0x00))
    else:
        set_font(run, 11, color=BLACK)

def add_divider(doc):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after  = Pt(4)
    run = p.add_run("─" * 80)
    set_font(run, 7, color=RGBColor(0xCC, 0xCC, 0xCC))

# ═════════════════════════════════════════════════════════════════════════════
# Header block
# ═════════════════════════════════════════════════════════════════════════════
p_title = doc.add_paragraph()
p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
p_title.paragraph_format.space_after = Pt(2)
shade_paragraph(p_title, 'FF9933')
r = p_title.add_run("  BharatDrishti — India's Eye  ")
set_font(r, 22, bold=True, color=RGBColor(0xFF,0xFF,0xFF), name="Calibri")

p_sub = doc.add_paragraph()
p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
p_sub.paragraph_format.space_before = Pt(0)
p_sub.paragraph_format.space_after  = Pt(2)
shade_paragraph(p_sub, '1A1A2E')
r2 = p_sub.add_run("  Team / Developer Profile  ·  Ideas of India Hackathon  ")
set_font(r2, 12, bold=False, color=RGBColor(0xCC,0xCC,0xCC), name="Calibri")

p_solo = doc.add_paragraph()
p_solo.alignment = WD_ALIGN_PARAGRAPH.CENTER
p_solo.paragraph_format.space_before = Pt(4)
p_solo.paragraph_format.space_after  = Pt(8)
r3 = p_solo.add_run("SOLO SUBMISSION  —  Entire project built by one developer")
set_font(r3, 11, bold=True, color=SAFFRON)

# ═════════════════════════════════════════════════════════════════════════════
# Section 1 — Project Details
# ═════════════════════════════════════════════════════════════════════════════
add_section_heading(doc, "1.  PROJECT DETAILS")
add_field(doc, "Project Title", "BharatDrishti — India's Eye")
add_field(doc, "Track",         "Sovereign Technology for India")
add_field(doc, "Category",      "AI Infra and Compute")
add_field(doc, "Hackathon",     "Ideas of India")
add_field(doc, "Submission Type", "Solo (single developer)")

# ═════════════════════════════════════════════════════════════════════════════
# Section 2 — Developer Profile
# ═════════════════════════════════════════════════════════════════════════════
add_section_heading(doc, "2.  DEVELOPER PROFILE")
add_field(doc, "Full Name",          "[YOUR NAME]",          placeholder=True)
add_field(doc, "Age",                "[YOUR AGE]",           placeholder=True)
add_field(doc, "College/Institution","[YOUR COLLEGE]",       placeholder=True)
add_field(doc, "Degree & Year",      "[YOUR DEGREE, YEAR]",  placeholder=True)
add_field(doc, "City",               "[YOUR CITY]",          placeholder=True)
add_field(doc, "Email",              "[YOUR EMAIL]",         placeholder=True)
add_field(doc, "Phone",              "[YOUR PHONE]",         placeholder=True)

# ═════════════════════════════════════════════════════════════════════════════
# Section 3 — Technical Skills
# ═════════════════════════════════════════════════════════════════════════════
add_section_heading(doc, "3.  TECHNICAL SKILLS USED IN THIS PROJECT")

skills = [
    ("Languages & Frameworks",
     ["Python 3.8", "JavaScript (ES2022)", "HTML / CSS"]),
    ("AI / ML",
     ["PyTorch", "ChangeFormerV6 Siamese Transformer (41M params)", "Segformer MiT-b2 encoder"]),
    ("Geospatial",
     ["Rasterio", "GDAL", "Shapely", "Fiona", "GeoJSON", "Sentinel-2 / Copernicus CDSE API"]),
    ("Backend",
     ["FastAPI", "Uvicorn", "Python-multiprocessing"]),
    ("Frontend",
     ["React 18", "Leaflet.js / react-leaflet", "react-compare-slider", "Vite"]),
    ("DevOps / Infrastructure",
     ["Conda environment", "Offline-first deployment", "ISRO / NIC cloud ready"]),
]

for category, items in skills:
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(5)
    p.paragraph_format.space_after  = Pt(1)
    p.paragraph_format.left_indent  = Inches(0.2)
    r = p.add_run(category)
    set_font(r, 11, bold=True, color=DGREY)
    for item in items:
        add_bullet(doc, item)

# ═════════════════════════════════════════════════════════════════════════════
# Section 4 — Solo Developer Role
# ═════════════════════════════════════════════════════════════════════════════
add_section_heading(doc, "4.  SOLO DEVELOPER ROLE")

p_intro = doc.add_paragraph()
p_intro.paragraph_format.space_before = Pt(4)
p_intro.paragraph_format.space_after  = Pt(4)
p_intro.paragraph_format.left_indent  = Inches(0.2)
r_intro = p_intro.add_run(
    "Built the entire BharatDrishti stack independently — "
    "from satellite data acquisition to the AI inference engine, "
    "REST API, React frontend, and all application documentation."
)
set_font(r_intro, 11, color=BLACK)

responsibilities = [
    ("Satellite Data Pipeline",
     "sentinel_fetch.py — CDSE OData integration, on-disk caching, synthetic GeoTIFF fallback"),
    ("AI Inference Engine",
     "inference.py — ChangeFormerV6 forward pass, NDVI computation, morphological mask cleaning, "
     "GeoJSON polygon extraction, spectral-diff fallback, natural-language AI brief generation"),
    ("REST API Backend",
     "api.py — FastAPI application with 10 endpoints, lazy model loading, response caching"),
    ("React Frontend",
     "frontend/ — before/after compare slider, Leaflet change map, NDVI diff panel, AI analysis panel, "
     "polygon statistics table"),
    ("Application Documentation",
     "README.md, application_answers.md — complete technical write-up, demo script, "
     "AI/ML deep-dive, limitations analysis"),
]

for role, detail in responsibilities:
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after  = Pt(1)
    p.paragraph_format.left_indent  = Inches(0.2)
    r_role = p.add_run(f"{role}:  ")
    set_font(r_role, 11, bold=True, color=SAFFRON)
    r_detail = p.add_run(detail)
    set_font(r_detail, 11, color=BLACK)

# ═════════════════════════════════════════════════════════════════════════════
# Section 5 — Brief Bio
# ═════════════════════════════════════════════════════════════════════════════
add_section_heading(doc, "5.  BRIEF BIO")

p_bio = doc.add_paragraph()
p_bio.paragraph_format.space_before = Pt(6)
p_bio.paragraph_format.space_after  = Pt(6)
p_bio.paragraph_format.left_indent  = Inches(0.2)
r_bio = p_bio.add_run(
    "[Write 2-3 lines about yourself here — your background, "
    "what drives you to build for India, and your interest in AI / geospatial technology.]"
)
set_font(r_bio, 11, color=RGBColor(0xAA, 0x66, 0x00))

# ═════════════════════════════════════════════════════════════════════════════
# Footer note
# ═════════════════════════════════════════════════════════════════════════════
add_divider(doc)
p_footer = doc.add_paragraph()
p_footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
r_f = p_footer.add_run(
    "BharatDrishti — built for Bharat, powered by open science.  "
    "·  Track: Sovereign Technology for India  ·  Ideas of India Hackathon"
)
set_font(r_f, 9, color=DGREY)

# ── Save ──────────────────────────────────────────────────────────────────────
out = r"C:\Users\dell\Desktop\times\ChangeFormer\team_profile.docx"
doc.save(out)
print(f"Saved: {out}")
